from datetime import date, datetime, timedelta, timezone
from io import BytesIO, StringIO
from pathlib import Path
import csv
import json
import os
import re
import shutil
import unicodedata
import uuid

from fastapi import FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import load_workbook
from PIL import Image, UnidentifiedImageError
from pydantic import BaseModel, Field, field_validator

from .documents import check_zip, convert_pdf, inspect_template, office_binary, render_batch
from .access import authorized, configured
from .defaults import seed_templates

ROOT = Path(__file__).resolve().parents[1]
DATA = Path(os.environ.get('CERTICOEX_DATA', ROOT / '.local/data'))
for folder in ('templates', 'jobs'):
    (DATA / folder).mkdir(parents=True, exist_ok=True)
seed_templates(DATA, ROOT / 'defaults')
app = FastAPI(title='Certicoex', docs_url=None, redoc_url=None)


@app.middleware('http')
async def local_requests(request: Request, call_next):
    if request.url.path != '/api/health':
        if not configured():
            return JSONResponse({'detail': 'Configura una contraseña de acceso de al menos 16 caracteres en el servicio.'}, status_code=503, headers={'Cache-Control':'no-store'})
        if not authorized(request.headers.get('authorization')):
            return JSONResponse({'detail': 'Introduce el usuario y la contraseña de la aplicación.'}, status_code=401, headers={'WWW-Authenticate':'Basic realm="Certicoex", charset="UTF-8"', 'Cache-Control':'no-store'})
    if request.method not in ('GET', 'HEAD', 'OPTIONS'):
        from urllib.parse import urlsplit
        origin = request.headers.get('origin')
        if origin and urlsplit(origin).netloc != request.headers.get('host'):
            return JSONResponse({'detail': 'Origen de solicitud no permitido.'}, status_code=403)
        if int(request.headers.get('content-length', 0)) > 15_000_000:
            return JSONResponse({'detail': 'El archivo supera el límite de 15 MB.'}, status_code=413)
    response = await call_next(request)
    response.headers['X-Content-Type-Options'] = 'nosniff'
    response.headers['Referrer-Policy'] = 'same-origin'
    response.headers['Cache-Control'] = 'no-store'
    response.headers['Content-Security-Policy'] = "default-src 'self'; script-src 'self'; style-src 'self'; img-src 'self' blob: data:; frame-src 'self'; object-src 'none'; base-uri 'self'; frame-ancestors 'self'"
    return response


@app.exception_handler(ValueError)
async def invalid_input(request, exc):
    return JSONResponse({'detail': str(exc)}, status_code=422)


def atomic_json(path, data):
    temp = path.with_suffix('.tmp')
    temp.write_text(json.dumps(data, ensure_ascii=False), encoding='utf8')
    temp.replace(path)


def template_path(key):
    if not re.fullmatch(r'[a-f0-9]{32}', key):
        raise HTTPException(404, 'Plantilla no encontrada.')
    path = DATA / 'templates' / (key + '.docx')
    if not path.is_file():
        raise HTTPException(404, 'Plantilla no encontrada.')
    return path


def store_template(data, name):
    fields = inspect_template(data)
    key = uuid.uuid4().hex
    (DATA / 'templates' / (key + '.docx')).write_bytes(data)
    meta = {'id': key, 'name': name.strip()[:120] or 'Curso sin título', 'fields': fields}
    atomic_json(DATA / 'templates' / (key + '.json'), meta)
    return meta


@app.get('/api/health')
def health():
    if not configured():
        raise HTTPException(503, 'Falta configurar el acceso privado.')
    return {'status': 'ok', 'pdf_available': bool(office_binary()), 'storage_ephemeral':os.environ.get('CERTICOEX_EPHEMERAL', '').lower() == 'true'}


@app.get('/api/templates')
def templates():
    return [json.loads(p.read_text(encoding='utf8')) for p in sorted((DATA / 'templates').glob('*.json'))]


@app.post('/api/templates')
def upload_template(file: UploadFile = File(...), name: str = Form(...)):
    if not (file.filename or '').lower().endswith('.docx'):
        raise ValueError('Sube una plantilla Word .docx.')
    data = file.file.read(10_000_001)
    if len(data) > 10_000_000:
        raise ValueError('La plantilla supera 10 MB.')
    try:
        return store_template(data, name)
    except (KeyError, TypeError) as exc:
        raise ValueError('No se puede leer esta plantilla Word.') from exc


@app.get('/api/profile')
def profile():
    path = DATA / 'profile.json'
    data = json.loads(path.read_text(encoding='utf8')) if path.exists() else {'name': ''}
    return {**data, 'has_signature': (DATA / 'signature.png').exists()}


@app.get('/api/profile/signature')
def signature_file():
    path = DATA / 'signature.png'
    if not path.exists():
        raise HTTPException(404)
    return FileResponse(path)


@app.post('/api/profile')
def save_profile(name: str = Form(...), signature: UploadFile | None = File(None)):
    name = name.strip()
    if not name or len(name) > 160:
        raise ValueError('Introduce tu nombre completo (máximo 160 caracteres).')
    if signature and signature.filename:
        raw = signature.file.read(5_000_001)
        if len(raw) > 5_000_000:
            raise ValueError('La firma debe ocupar menos de 5 MB.')
        try:
            with Image.open(BytesIO(raw)) as im:
                if im.format not in ('PNG', 'JPEG') or im.width * im.height > 12_000_000:
                    raise ValueError('Usa una firma PNG o JPG de hasta 12 megapíxeles.')
                im.load()
                out = BytesIO()
                im.convert('RGBA').save(out, format='PNG')
                temp = DATA / 'signature.tmp'
                temp.write_bytes(out.getvalue())
                temp.replace(DATA / 'signature.png')
        except (UnidentifiedImageError, OSError, Image.DecompressionBombError) as exc:
            raise ValueError('No se puede leer la imagen de firma.') from exc
    atomic_json(DATA / 'profile.json', {'name': name})
    return profile()


def normalize(value):
    return ''.join(c for c in unicodedata.normalize('NFKD', str(value).lower()) if not unicodedata.combining(c)).strip()


def tabular(rows):
    rows = [list(row) for row in rows if any(str(c or '').strip() for c in row)]
    if not rows:
        raise ValueError('El listado está vacío.')
    if len(rows) > 501 or max(map(len, rows)) > 40:
        raise ValueError('Admite hasta 500 alumnos y 40 columnas por importación.')
    aliases = {'name': ['nombre completo', 'nombre y apellidos', 'trabajador', 'alumno', 'nombre', 'nombres'], 'surname': ['apellidos', 'apellido', 'primer apellido'], 'surname2': ['segundo apellido'], 'dni': ['dni/nie', 'dni', 'nie', 'nif', 'documento', 'dni / nie']}
    first = [normalize(c or '') for c in rows[0]]
    mapping = {field: next((first.index(a) for a in names if a in first), -1) for field, names in aliases.items()}
    has_header = mapping['name'] >= 0 or mapping['dni'] >= 0
    width = max(map(len, rows))
    headers = [str(c or f'Columna {i+1}') for i, c in enumerate(rows[0])] if has_header else [f'Columna {i+1}' for i in range(width)]
    if not has_header:
        mapping = {'name': 0, 'surname': -1, 'surname2': -1, 'dni': 1 if width > 1 else -1}
    return {'headers': headers, 'rows': [[str(c).strip() if c is not None else '' for c in row] for row in rows[int(has_header):]], 'mapping': mapping}


class Paste(BaseModel):
    text: str = Field(max_length=150_000)


@app.post('/api/import/paste')
def paste_list(payload: Paste):
    text = payload.text.strip()
    if not text:
        raise ValueError('Pega el listado de alumnos.')
    try:
        dialect = csv.Sniffer().sniff(text[:4000], delimiters='\t;,')
        rows = list(csv.reader(StringIO(text), dialect))
    except csv.Error:
        # Also accept "Nombre Apellidos 12345678Z" on each line.
        rows = []
        for line in text.splitlines():
            match = re.match(r'^(.*?)\s+([XYZxyz\d][\d .-]{5,}[A-Za-z])\s*$', line)
            rows.append([match[1], match[2]] if match else [line, ''])
    return tabular(rows)


@app.post('/api/import/excel')
def import_excel(file: UploadFile = File(...)):
    if not (file.filename or '').lower().endswith('.xlsx'):
        raise ValueError('Usa Excel .xlsx. Para .xls, guárdalo como .xlsx o copia y pega las celdas.')
    raw = file.file.read(10_000_001)
    if len(raw) > 10_000_000:
        raise ValueError('El Excel supera 10 MB.')
    check_zip(raw)
    try:
        workbook = load_workbook(BytesIO(raw), read_only=True, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(min_row=1, max_row=502, max_col=41, values_only=True))
        # Trim unused trailing columns, preserving interior empty cells.
        rows = [list(row) for row in rows]
        for row in rows:
            while row and row[-1] is None:
                row.pop()
        result = tabular(rows)
        result['sheet'] = sheet.title
        workbook.close()
        return result
    except (KeyError, TypeError, OSError) as exc:
        raise ValueError('No se puede leer el Excel. Comprueba que es un archivo .xlsx válido.') from exc


class Student(BaseModel):
    name: str = Field(min_length=1, max_length=180)
    dni: str = Field(min_length=1, max_length=30)

    @field_validator('name', 'dni')
    @classmethod
    def trim(cls, value):
        value = value.strip()
        if not value or any(ord(c) < 32 for c in value):
            raise ValueError('Nombre y documento deben contener texto válido.')
        return value


class Batch(BaseModel):
    template_id: str
    company: str = Field(min_length=1, max_length=200)
    cif: str = Field(min_length=1, max_length=30)
    hours: float = Field(gt=0, le=1000, allow_inf_nan=False)
    course_date: date
    students: list[Student] = Field(min_length=1, max_length=500)
    preview: bool = False

    @field_validator('company', 'cif')
    @classmethod
    def nonblank(cls, value):
        value = value.strip()
        if not value or any(ord(c) < 32 for c in value):
            raise ValueError('Completa los datos de la empresa.')
        return value


MONTHS = ['enero', 'febrero', 'marzo', 'abril', 'mayo', 'junio', 'julio', 'agosto', 'septiembre', 'octubre', 'noviembre', 'diciembre']


def clean_jobs():
    cutoff = datetime.now(timezone.utc).timestamp() - timedelta(hours=24).total_seconds()
    for folder in (DATA / 'jobs').iterdir():
        if folder.is_dir() and folder.stat().st_mtime < cutoff:
            shutil.rmtree(folder)


@app.post('/api/generate')
def generate(payload: Batch):
    template = template_path(payload.template_id)
    current = profile()
    if not current['name'] or not current['has_signature']:
        raise ValueError('Guarda tu nombre y una imagen de firma en el perfil del técnico.')
    seen = set()
    for s in payload.students:
        dni = re.sub(r'[\s.\-]', '', s.dni).upper()
        if dni in seen:
            raise ValueError('Hay documentos de identidad repetidos. Corrige el listado antes de generar.')
        seen.add(dni)
    clean_jobs()
    key = uuid.uuid4().hex
    folder = DATA / 'jobs' / key
    folder.mkdir()
    d = payload.course_date
    common = {'EMPRESA': payload.company, 'CIF': payload.cif, 'HORAS': f'{payload.hours:g}', 'FECHA': f'{d.day} de {MONTHS[d.month-1]} de {d.year}', 'TECNICO': current['name']}
    students = [s.model_dump() for s in (payload.students[:1] if payload.preview else payload.students)]
    output = folder / 'certificados.docx'
    render_batch(template, students, common, (DATA / 'signature.png').read_bytes(), output)
    result = {'word': f'/api/files/{key}/docx', 'count': len(students), 'pdf': None, 'pages': None}
    try:
        _, pages = convert_pdf(output, folder)
        result.update(pdf=f'/api/files/{key}/pdf', pages=pages)
    except RuntimeError as exc:
        result['warning'] = str(exc)
    return result


@app.get('/api/files/{key}/{kind}')
def download(key: str, kind: str, inline: bool = False):
    if not re.fullmatch(r'[a-f0-9]{32}', key) or kind not in ('pdf', 'docx'):
        raise HTTPException(404)
    path = DATA / 'jobs' / key / ('certificados.' + kind)
    if not path.exists() or path.stat().st_mtime < datetime.now(timezone.utc).timestamp() - 86400:
        raise HTTPException(404, 'Archivo caducado. Vuelve a generar los certificados.')
    return FileResponse(path, media_type='application/pdf' if kind == 'pdf' else 'application/vnd.openxmlformats-officedocument.wordprocessingml.document', filename=path.name, content_disposition_type='inline' if inline else 'attachment')


app.mount('/', StaticFiles(directory=ROOT / 'static', html=True), name='web')
