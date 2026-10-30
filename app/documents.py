"""Fill OOXML runs without flattening the uploaded document's formatting."""
from copy import deepcopy
from io import BytesIO
from pathlib import Path
import re
import shutil
import subprocess
import os
from zipfile import ZipFile, BadZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm
from PIL import Image

FIELDS = {'NOMBRE', 'DNI/NIE', 'EMPRESA', 'CIF', 'HORAS', 'FECHA', 'TECNICO', 'FIRMA'}
SIGNATURE_MARKER = 'certicoex_signature'


def check_zip(data: bytes):
    try:
        with ZipFile(BytesIO(data)) as archive:
            infos = archive.infolist()
            if len(infos) > 3000 or sum(i.file_size for i in infos) > 60_000_000:
                raise ValueError('El archivo descomprimido es demasiado grande.')
            if any('vbaProject' in i.filename for i in infos):
                raise ValueError('No se admiten documentos con macros.')
    except BadZipFile as exc:
        raise ValueError('El archivo no es un documento Office válido.') from exc


def paragraphs(root):
    return root.iter(qn('w:p'))


def paragraph_text(p):
    return ''.join(t.text or '' for t in p.iter(qn('w:t')))


def replace_text(p, values):
    # Replace backwards, retaining the formatting of the first character even
    # when Word has split a placeholder across multiple runs.
    nodes = list(p.iter(qn('w:t')))
    original = ''.join(n.text or '' for n in nodes)
    spans = []
    pos = 0
    for n in nodes:
        end = pos + len(n.text or '')
        spans.append((n, pos, end))
        pos = end
    for match in reversed(list(re.finditer(r'\{([^{}]+)\}', original))):
        if match.group(1) not in values:
            continue
        first = True
        for n, start, end in spans:
            if end <= match.start() or start >= match.end():
                continue
            a, b = max(0, match.start()-start), min(end-start, match.end()-start)
            text = n.text or ''
            n.text = text[:a] + (values[match.group(1)] if first else '') + text[b:]
            n.set(qn('xml:space'), 'preserve')
            first = False


def document_roots(doc):
    roots = [doc.element]
    # Include all header/footer parts, including inherited headers.
    for rel in doc.part.rels.values():
        if not rel.is_external and rel.reltype.rsplit('/', 1)[-1] in ('header', 'footer'):
            roots.append(rel.target_part.element)
    return roots


def inspect_template(data):
    check_zip(data)
    doc = Document(BytesIO(data))
    roots = document_roots(doc)
    fields = set()
    for root in roots:
        for p in paragraphs(root):
            fields.update(re.findall(r'\{([^{}]+)\}', paragraph_text(p)))
    missing = {'NOMBRE', 'DNI/NIE', 'EMPRESA', 'CIF', 'HORAS', 'FECHA', 'TECNICO'} - fields
    anchored = any(n.get('descr') == SIGNATURE_MARKER for n in doc.element.iter(qn('wp:docPr')))
    if not anchored and 'FIRMA' not in fields:
        missing.add('FIRMA')
    if missing:
        raise ValueError('Faltan campos en la plantilla: ' + ', '.join(sorted(missing)))
    unknown = fields - FIELDS
    if unknown:
        raise ValueError('Campos no reconocidos: ' + ', '.join(sorted(unknown)))
    for root in roots[1:]:
        for p in paragraphs(root):
            if re.search(r'\{(?:NOMBRE|DNI/NIE|FIRMA)\}', paragraph_text(p)):
                raise ValueError('Coloca los campos de alumno y firma en el cuerpo, fuera de cabeceras y pies.')
    return sorted(fields)


def prepare_example(source: Path, destination: Path):
    """Adapt the provided example locally; never modify the original upload."""
    doc = Document(source)
    ps = list(paragraphs(doc.element))
    for index, p in enumerate(ps):
        if paragraph_text(p).strip() == 'TÉCNICO EN PRL':
            for candidate in ps[index+1:]:
                if paragraph_text(candidate).strip():
                    text = paragraph_text(candidate)
                    # Preserve paragraph/run properties when replacing the fixed name.
                    nodes = list(candidate.iter(qn('w:t')))
                    nodes[0].text = '{TECNICO}'
                    for node in nodes[1:]:
                        node.text = ''
                    break
            break
    drawings = list(doc.element.iter(qn('w:drawing')))
    if len(drawings) != 1:
        raise ValueError('La adaptación de ejemplo requiere una sola firma en el cuerpo.')
    drawing = drawings[0]
    blank = BytesIO()
    Image.new('RGBA', (200, 80), (255, 255, 255, 0)).save(blank, format='PNG')
    rid, _ = doc.part.get_or_add_image(BytesIO(blank.getvalue()))
    old_rids = set()
    for blip in drawing.iter(qn('a:blip')):
        old_rids.add(blip.get(qn('r:embed')))
        blip.set(qn('r:embed'), rid)
    for pr in drawing.iter(qn('wp:docPr')):
        pr.set('descr', SIGNATURE_MARKER)
    for old in old_rids:
        if old and old != rid and not doc.element.xpath('//*[@r:embed="' + old + '"]'):
            doc.part.drop_rel(old)
    doc.save(destination)
    inspect_template(destination.read_bytes())


def put_signature(doc, signature):
    rid, _ = doc.part.get_or_add_image(BytesIO(signature))
    with Image.open(BytesIO(signature)) as im:
        ratio = im.width / im.height
    for drawing in doc.element.iter(qn('w:drawing')):
        markers = list(drawing.iter(qn('wp:docPr')))
        if any(m.get('descr') == SIGNATURE_MARKER for m in markers):
            for blip in drawing.iter(qn('a:blip')):
                blip.set(qn('r:embed'), rid)
            extent = next(drawing.iter(qn('wp:extent')), None)
            if extent is not None:
                width = int(extent.get('cx'))
                height = int(extent.get('cy'))
                width = min(width, int(height * ratio))
                height = int(width / ratio)
                extent.set('cx', str(width))
                extent.set('cy', str(height))
                for ex in drawing.iter(qn('a:ext')):
                    ex.set('cx', str(width))
                    ex.set('cy', str(height))
    # Generic uploaded templates can place {FIRMA} in a standalone paragraph.
    from docx.text.paragraph import Paragraph
    for p in paragraphs(doc.element):
        if '{FIRMA}' in paragraph_text(p):
            replace_text(p, {'FIRMA': ''})
            Paragraph(p, doc).add_run().add_picture(BytesIO(signature), width=Mm(35))


def render_batch(template: Path, students, common, signature: bytes, output: Path):
    doc = Document(template)
    put_signature(doc, signature)
    for root in document_roots(doc)[1:]:
        for p in paragraphs(root):
            replace_text(p, common)
    body = doc.element.body
    final_section = deepcopy(body.sectPr)
    content = [deepcopy(el) for el in body if el.tag != qn('w:sectPr')]
    styles = {s.get(qn('w:styleId')): s for s in doc.styles.element}
    for element in content:
        for p in paragraphs(element):
            props = p.find(qn('w:pPr'))
            if props is None or props.find(qn('w:numPr')) is not None:
                continue
            ref = props.find(qn('w:pStyle'))
            style_id = ref.get(qn('w:val')) if ref is not None else None
            visited = set()
            while style_id in styles and style_id not in visited:
                visited.add(style_id)
                style = styles[style_id]
                num_props = style.find('./' + qn('w:pPr') + '/' + qn('w:numPr'))
                if num_props is not None:
                    props.append(deepcopy(num_props))
                    break
                parent = style.find(qn('w:basedOn'))
                style_id = parent.get(qn('w:val')) if parent is not None else None
    numbering = doc.part.numbering_part.element
    original_nums = list(numbering.findall(qn('w:num')))
    original_abstracts = {el.get(qn('w:abstractNumId')): el for el in numbering.findall(qn('w:abstractNum'))}
    next_num = max([int(el.get(qn('w:numId'))) for el in original_nums] + [0]) + 1
    next_abstract = max([int(i) for i in original_abstracts] + [0]) + 1
    # A section boundary retains the final columns/header of the certificate;
    # the next certificate's first section starts on a fresh page.
    for el in list(body):
        body.remove(el)
    for index, student in enumerate(students):
        values = {**common, 'NOMBRE': student['name'], 'DNI/NIE': student['dni']}
        num_map = {}
        # Independent numbering definitions prevent the next student's temario
        # from continuing at 10, 11, ... in either Word or LibreOffice.
        for original in original_nums:
            num = deepcopy(original)
            old_id = original.get(qn('w:numId'))
            num_map[old_id] = str(next_num)
            num.set(qn('w:numId'), str(next_num))
            for attribute in num.attrib:
                if attribute.endswith('}durableId'):
                    num.set(attribute, str(next_num))
            abstract_ref = num.find(qn('w:abstractNumId'))
            abstract = deepcopy(original_abstracts[abstract_ref.get(qn('w:val'))])
            abstract.set(qn('w:abstractNumId'), str(next_abstract))
            nsid = abstract.find(qn('w:nsid'))
            if nsid is not None:
                nsid.set(qn('w:val'), f'{next_abstract:08X}')
            abstract_ref.set(qn('w:val'), str(next_abstract))
            # abstractNum must precede num elements in numbering.xml.
            first_num = numbering.find(qn('w:num'))
            numbering.insert(numbering.index(first_num), abstract)
            numbering.append(num)
            next_num += 1
            next_abstract += 1
        if index:
            boundary = OxmlElement('w:p')
            props = OxmlElement('w:pPr')
            props.append(deepcopy(final_section))
            boundary.append(props)
            body.append(boundary)
        for element in content:
            clone = deepcopy(element)
            for num_id in clone.iter(qn('w:numId')):
                if num_id.get(qn('w:val')) in num_map:
                    num_id.set(qn('w:val'), num_map[num_id.get(qn('w:val'))])
            for p in paragraphs(clone):
                replace_text(p, values)
            body.append(clone)
    body.append(final_section)
    for index, pr in enumerate(doc.element.iter(qn('wp:docPr')), 1):
        pr.set('id', str(index))
    doc.save(output)


def office_binary():
    configured = os.environ.get('CERTICOEX_LIBREOFFICE')
    if configured:
        return configured
    local = Path(__file__).resolve().parents[1] / '.local/office/usr/lib/libreoffice/program/soffice'
    found = shutil.which('libreoffice') or shutil.which('soffice')
    if found:
        return found
    candidates = [local, Path('/Applications/LibreOffice.app/Contents/MacOS/soffice')]
    for key in ('PROGRAMFILES', 'PROGRAMFILES(X86)'):
        base = os.environ.get(key)
        if base:
            candidates.append(Path(base) / 'LibreOffice/program/soffice.exe')
    return next((str(p) for p in candidates if p.is_file()), None)


def convert_pdf(docx: Path, outdir: Path):
    binary = office_binary()
    if not binary:
        raise RuntimeError('Falta LibreOffice para generar el PDF. El Word sí está disponible.')
    import tempfile
    with tempfile.TemporaryDirectory(prefix='certicoex-office-') as profile:
        env = os.environ.copy()
        local = Path(__file__).resolve().parents[1] / '.local/office'
        if str(local) in binary:
            libs = [local / 'usr/lib/libreoffice/program', local / 'usr/lib/x86_64-linux-gnu']
            env['LD_LIBRARY_PATH'] = ':'.join(str(p) for p in libs) + ':' + env.get('LD_LIBRARY_PATH', '')
            env['FONTCONFIG_FILE'] = str(local / 'fonts.conf')
        env['SAL_USE_VCLPLUGIN'] = 'svp'
        env['XDG_CACHE_HOME'] = str(Path(profile) / 'cache')
        try:
            result = subprocess.run([binary, '-env:UserInstallation=' + Path(profile).as_uri(), '--headless', '--convert-to', 'pdf:writer_pdf_Export', '--outdir', str(outdir), str(docx)], env=env, capture_output=True, timeout=180)
        except subprocess.TimeoutExpired as exc:
            raise RuntimeError('La conversión a PDF superó el tiempo disponible. Puedes descargar el Word y volver a intentarlo.') from exc
    target = outdir / (docx.stem + '.pdf')
    if result.returncode or not target.exists():
        raise RuntimeError('LibreOffice no pudo convertir el documento a PDF. El Word sigue disponible.')
    from pypdf import PdfReader
    reader = PdfReader(target)
    if not reader.pages:
        raise RuntimeError('El PDF generado está vacío.')
    return target, len(reader.pages)
