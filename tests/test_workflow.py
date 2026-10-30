import os
import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import patch
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from fastapi.testclient import TestClient
from openpyxl import Workbook
from PIL import Image
from pypdf import PdfReader
from app import main
from app.documents import render_batch


def signature():
    b = BytesIO()
    Image.new('RGBA', (300, 100), (20, 100, 80, 255)).save(b, format='PNG')
    return b.getvalue()


class Workflow(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.data = Path(self.temp.name)
        for folder in ('templates', 'jobs'):
            (self.data / folder).mkdir()
        self.patcher = patch.object(main, 'DATA', self.data)
        self.patcher.start()
        self.client = TestClient(main.app)
        doc = Document()
        doc.sections[0].header.paragraphs[0].text = 'Cabecera · {EMPRESA}'
        p = doc.add_paragraph()
        p.add_run('{NOM').bold = True
        p.add_run('BRE}')
        doc.add_paragraph('{DNI/NIE} · {EMPRESA} · {CIF} · {HORAS} · {FECHA}')
        doc.add_paragraph('{FIRMA}')
        doc.add_paragraph('{TECNICO}')
        numbering = doc.part.numbering_part.element
        abstract = OxmlElement('w:abstractNum'); abstract.set(qn('w:abstractNumId'), '100')
        level = OxmlElement('w:lvl'); level.set(qn('w:ilvl'), '0')
        for tag, value in [('start','1'), ('numFmt','decimal'), ('lvlText','%1.')]:
            element=OxmlElement('w:'+tag); element.set(qn('w:val'),value); level.append(element)
        abstract.append(level); numbering.insert(0, abstract)
        num=OxmlElement('w:num'); num.set(qn('w:numId'),'100')
        ref=OxmlElement('w:abstractNumId'); ref.set(qn('w:val'),'100'); num.append(ref); numbering.append(num)
        p=doc.add_paragraph('Tema de prueba')
        props=OxmlElement('w:numPr'); number=OxmlElement('w:numId'); number.set(qn('w:val'),'100'); props.append(number)
        style=doc.styles.add_style('Temario de prueba',WD_STYLE_TYPE.PARAGRAPH)
        style.element.get_or_add_pPr().append(props)
        p.style=style
        b = BytesIO(); doc.save(b)
        result = self.client.post('/api/templates', data={'name':'Curso de prueba'}, files={'file':('plantilla.docx',b.getvalue())})
        self.assertEqual(result.status_code, 200, result.text)
        self.template = result.json()['id']
        result = self.client.post('/api/profile', data={'name':'Técnico de prueba'}, files={'signature':('firma.png',signature(),'image/png')})
        self.assertEqual(result.status_code, 200, result.text)

    def tearDown(self):
        self.client.close(); self.patcher.stop(); self.temp.cleanup()

    def batch(self):
        return {'template_id':self.template, 'company':'Empresa & Hijos <Prueba>', 'cif':'B12345678','hours':8,'course_date':'2026-10-30','students':[{'name':'Ana García López','dni':'12345678Z'},{'name':'Luis Pérez Martín','dni':'87654321X'}]}

    def test_split_runs_formatting_and_xml_escaping(self):
        output = self.data/'render.docx'
        batch = self.batch()
        common = {'EMPRESA':batch['company'],'CIF':batch['cif'],'HORAS':'8','FECHA':'30 de octubre de 2026','TECNICO':'Técnico de prueba'}
        render_batch(main.template_path(self.template),batch['students'],common,signature(),output)
        doc=Document(output)
        text='\n'.join(p.text for p in doc.paragraphs)
        for student in batch['students']:
            self.assertEqual(text.count(student['name']),1)
        self.assertNotIn('{',text)
        self.assertIn(batch['company'],text)
        self.assertTrue(doc.paragraphs[0].runs[0].bold)
        self.assertIn(batch['company'],doc.sections[0].header.paragraphs[0].text)
        self.assertEqual(len(doc.inline_shapes),2)

    def test_paste_and_excel_column_detection(self):
        pasted = self.client.post('/api/import/paste',json={'text':'Nombre;Apellidos;DNI\nAna;García López;12345678Z\nLuis;Pérez;87654321X'})
        self.assertEqual(pasted.status_code,200)
        self.assertEqual(pasted.json()['mapping'],{'name':0,'surname':1,'surname2':-1,'dni':2})
        w=Workbook();s=w.active;s.append(['Nombre completo','DNI/NIE'])
        for i in range(50):s.append([f'Alumno {i}',f'{i:08d}A'])
        b=BytesIO();w.save(b)
        response=self.client.post('/api/import/excel',files={'file':('alumnos.xlsx',b.getvalue())})
        self.assertEqual(response.status_code,200,response.text)
        self.assertEqual(len(response.json()['rows']),50)
        self.assertEqual(response.json()['mapping']['dni'],1)
        self.assertEqual(response.json()['rows'][0][1],'00000000A')

    def test_validation_and_duplicate_identity(self):
        batch=self.batch();batch['students'][1]['dni']='12.345.678-z'
        self.assertEqual(self.client.post('/api/generate',json=batch).status_code,422)
        for key,value in [('company','  '),('hours',0)]:
            batch=self.batch();batch[key]=value
            self.assertEqual(self.client.post('/api/generate',json=batch).status_code,422)
        batch=self.batch();batch['students'][0]['name']='  '
        self.assertEqual(self.client.post('/api/generate',json=batch).status_code,422)

    def test_signature_required_and_malformed_files(self):
        (self.data/'signature.png').unlink()
        self.assertEqual(self.client.post('/api/generate',json=self.batch()).status_code,422)
        self.assertEqual(self.client.post('/api/import/excel',files={'file':('x.xlsx',b'not a zip')}).status_code,422)
        self.assertEqual(self.client.post('/api/profile',data={'name':'Test'},files={'signature':('x.png',b'not an image')}).status_code,422)
        self.assertEqual(self.client.get('/api/profile').json()['name'],'Técnico de prueba')

    def test_no_cross_origin_write(self):
        response=self.client.post('/api/import/paste',json={'text':'Ana;12345678Z'},headers={'Origin':'https://otro.example'})
        self.assertEqual(response.status_code,403)

    def test_word_available_when_converter_unavailable(self):
        with patch('app.main.convert_pdf',side_effect=RuntimeError('Falta LibreOffice')):
            response=self.client.post('/api/generate',json=self.batch())
        self.assertEqual(response.status_code,200,response.text)
        self.assertIsNone(response.json()['pdf'])
        self.assertEqual(self.client.get(response.json()['word']).status_code,200)

    @unittest.skipUnless(os.environ.get('CERTICOEX_TEST_PDF') == '1', 'Activar con CERTICOEX_TEST_PDF=1')
    def test_real_pdf_and_preview(self):
        for preview in (False, True):
            batch=self.batch();batch['preview']=preview
            response=self.client.post('/api/generate',json=batch)
            self.assertEqual(response.status_code,200,response.text)
            result=response.json();self.assertIsNotNone(result['pdf'], result)
            pdf=self.client.get(result['pdf']);self.assertEqual(pdf.status_code,200)
            reader=PdfReader(BytesIO(pdf.content))
            self.assertEqual(len(reader.pages),1 if preview else 2)
            text='\n'.join(p.extract_text() for p in reader.pages)
            self.assertIn('Ana García López',text)
            self.assertEqual('Luis Pérez Martín' in text,not preview)
            for page in reader.pages:
                self.assertRegex(page.extract_text(),r'1\.\s+Tema de prueba')


if __name__ == '__main__':
    unittest.main()
